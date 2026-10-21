"""Pruebas de TuitIA con Playwright: IA, servicios de tuits y API de capturas SIMULADOS (sin claves ni red reales).
       python3 build.py && python3 tests/smoke_test.py
   La IA simulada miente a propósito: una idea con una cita que no está en el guion, un enlace de tuit inventado,
   una frase clave que no está en el tuit, un id de tuit que no existe y una cifra nueva. El código debe detectarlo todo."""
import asyncio, io, json, pathlib, re, subprocess, sys, tempfile
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from forja_test import Harness
from PIL import Image

ROOT = pathlib.Path(__file__).resolve().parent.parent
TMP = pathlib.Path(tempfile.mkdtemp(prefix='tuitia_'))
SHOTS = ROOT / 'tests' / 'capturas'

def png(w, h, color=(255, 255, 255), caja=None):
    im = Image.new('RGB', (w, h), color)
    if caja:
        x0, y0, x1, y1, c = caja; im.paste(c, (x0, y0, x1, y1))
    b = io.BytesIO(); im.save(b, 'PNG'); return b.getvalue()

GUION = ('Clase 2. El cambio climático y la opinión pública.\n\n'
         'Primero, el consenso científico sobre el origen humano del calentamiento es amplísimo y se mantiene estable desde hace años.\n\n'
         'Segundo, la percepción del público no coincide con ese consenso: mucha gente cree que los científicos están divididos.\n\n'
         'Tercero, las olas de calor recientes han hecho que el tema aparezca más en las conversaciones cotidianas de las redes.')
# tuits simulados (FxTwitter); 1800000000000000222 no existe; 1800000000000000333 solo responde vxTwitter; 1800000000000000444 lo encuentra la búsqueda web; 1800000000000000555 lo inventa la IA
FX = {'1800000000000000111': {'tweet': {'text': 'El consenso científico sobre el cambio climático es sólido: más del 97 % de los estudios coinciden. https://t.co/abc', 'author': {'name': 'Ciencia Abierta', 'screen_name': 'ciencia_abierta', 'avatar_url': 'https://pbs.twimg.com/profile_images/1/a_normal.png'}, 'created_timestamp': 1760000000, 'likes': 1520, 'retweets': 300, 'replies': 41, 'views': 90000, 'media': {'photos': [{'url': 'https://pbs.twimg.com/media/foto.png'}]}}},
      '1800000000000000444': {'tweet': {'text': 'Otra ola de calor y otra vez el clima es lo único de lo que se habla en la terraza del bar.', 'author': {'name': 'Vecina Observadora', 'screen_name': 'vecina_obs', 'avatar_url': ''}, 'created_timestamp': 1761000000, 'likes': 88, 'retweets': 9, 'replies': 3}}}
VX = {'1800000000000000333': {'text': 'Encuesta: la mitad cree que los científicos no se ponen de acuerdo sobre el clima.', 'user_name': 'Datos y Opinión', 'user_screen_name': 'datos_op', 'user_profile_image_url': '', 'date_epoch': 1759000000, 'likes': 400, 'retweets': 120, 'replies': 15, 'media_extended': []}}

async def ia(prov, system, user, info):
    if 'Sepáralo en sus ideas' in user or 'Separate it' in user:
        return ('j', {'titulo': 'Clima y opinión pública', 'ideas': [
            {'titulo': 'Consenso científico', 'texto': 'El consenso sobre el origen humano es amplio y estable.', 'cita': 'el consenso científico sobre el origen humano del calentamiento es amplísimo', 'consultas': ['consenso científico clima lang:es min_faves:50', 'consenso científico clima lang:es min_faves:50', '"97 %" científicos clima lang:es']},
            {'titulo': 'Percepción dividida', 'texto': 'El público cree que hay división científica.', 'cita': 'mucha gente cree que los científicos están divididos', 'consultas': ['científicos divididos clima lang:es']},
            {'titulo': 'El calor en la conversación', 'texto': 'Las olas de calor llevan el tema a las redes en 2024.', 'cita': 'las olas de calor del verano pasado multiplicaron los tuits', 'consultas': ['ola de calor clima conversación lang:es']}]})
    if 'valora si refuerza' in user:
        return ('j', {'valoraciones': [
            {'id': '1800000000000000111', 'encaje': 'alto', 'porque': 'Afirma el consenso con una cifra.', 'frase': 'El consenso científico sobre el cambio climático es sólido', 'titulo': 'Un consenso sólido', 'nota': 'Comenta la cifra del 97 % con la clase.'},
            {'id': '1800000000000000333', 'encaje': 'medio', 'porque': 'Muestra la percepción de división.', 'frase': 'los científicos discrepan muchísimo entre ellos', 'titulo': 'La percepción', 'nota': 'Según el 2023 Survey, 65 % lo cree.'},
            {'id': '1800000000000000999', 'encaje': 'alto', 'porque': 'Tuit que no existe', 'frase': 'x', 'titulo': 'x', 'nota': 'x'}]})
    return ('t', '{}')

async def main():
    ok = True
    def check(nombre, cond, extra=''):
        nonlocal ok; ok &= bool(cond); print(('✓' if cond else '✕'), nombre, extra if not cond else '')
    SHOTS.mkdir(exist_ok=True)
    async with Harness(ROOT / 'index.html', ia) as h:
        p = h.page
        vistos = {'fx': [], 'vx': [], 'syn': [], 'cap': [], 'img': []}
        async def r_fx(r):
            i = r.request.url.rstrip('/').split('/')[-1]; vistos['fx'].append(i)
            if i in FX: return await r.fulfill(status=200, content_type='application/json', body=json.dumps(FX[i]), headers={'access-control-allow-origin': '*'})
            return await r.fulfill(status=404 if i != '1800000000000000333' else 500, body='{}', headers={'access-control-allow-origin': '*'})
        async def r_vx(r):
            i = r.request.url.rstrip('/').split('/')[-1]; vistos['vx'].append(i)
            if i in VX: return await r.fulfill(status=200, content_type='application/json', body=json.dumps(VX[i]), headers={'access-control-allow-origin': '*'})
            return await r.fulfill(status=404, body='{}', headers={'access-control-allow-origin': '*'})
        async def r_syn(r): vistos['syn'].append(r.request.url); await r.fulfill(status=404, body='{}', headers={'access-control-allow-origin': '*'})
        async def r_img(r): vistos['img'].append(r.request.url); await r.fulfill(status=200, content_type='image/png', body=png(400, 300, (30, 120, 200)), headers={'access-control-allow-origin': '*'})
        CAP = {'modo': 'ok'}
        async def r_cap(r):
            vistos['cap'].append(r.request.url)
            if CAP['modo'] == '429': return await r.fulfill(status=429, body='Too many requests', headers={'access-control-allow-origin': '*'})
            await r.fulfill(status=200, content_type='image/png', body=png(1200, 2200, (255, 255, 255), (100, 100, 1100, 700, (20, 20, 20))), headers={'access-control-allow-origin': '*'})
        await p.route('https://api.fxtwitter.com/**', r_fx); await p.route('https://api.vxtwitter.com/**', r_vx)
        await p.route('https://cdn.syndication.twimg.com/**', r_syn); await p.route('https://pbs.twimg.com/**', r_img)
        for host in ('https://api.microlink.io/**', 'https://api.screenshotone.com/**', 'https://api.apiflash.com/**'): await p.route(host, r_cap)

        check('CSP bloquea otros orígenes', await h.csp_bloquea())
        check('CSP permite solo los servicios declarados', 'https://api.fxtwitter.com' in await p.evaluate("document.querySelector('meta[http-equiv]').content"))
        # ── modelos e idiomas ──
        for prov, esperado in (('openai', 'gpt-5-mini'), ('gemini', 'gemini-2.5-pro'), ('anthropic', 'claude-sonnet-4-5')):
            ops = await h.modelos(prov)
            check(f'desplegable de modelos ({prov})', esperado in ops and '__otro' in ops, str(ops))
        await p.keyboard.press('Escape'); await p.evaluate("localStorage.setItem(APP+'_prov','openai')")
        check('no faltan traducciones', await h.faltan_traducciones() == [], str(await h.faltan_traducciones())[:300])
        es_btn = await p.inner_text('#btnStart')
        await h.idioma('en'); check('inglés: portada y barra', await p.inner_text('#btnStart') != es_btn and (await p.inner_text('#btnHelp')).strip().lower() == 'help')
        await h.idioma('pt'); check('portugués: portada y barra', (await p.inner_text('#btnHelp')).strip().lower() == 'ajuda')
        await h.idioma('es'); check('vuelta al español', await p.inner_text('#btnStart') == es_btn)
        await h.shot(str(SHOTS / 'portada.png'))
        # ── ejemplo sin clave, en cada idioma ──
        for l, palabra in (('es', 'Desinformación'), ('en', 'Misinformation'), ('pt', 'Desinformação')):
            await h.idioma(l); await p.evaluate('go(0)'); await p.click('#btnEjemplo'); await h.esp("step===3 && document.querySelectorAll('img[data-img]:not([hidden])').length>=4")
            r = await p.evaluate("({tit:S.titulo, ideas:S.ideas.every(i=>i.citaOk), frases:Object.values(S.tuits).every(fraseOk), n:Object.keys(S.tuits).length})")
            check(f'ejemplo en {l}: texto propio, citas del guion y frases clave verificadas, tarjetas dibujadas', palabra in r['tit'] and r['ideas'] and r['frases'] and r['n'] == 4, str(r))
        await h.idioma('en'); await p.wait_for_timeout(500)
        check('viendo el ejemplo, cambiar de idioma lo recarga en ese idioma', 'Misinformation' in await p.evaluate('S.titulo'))
        await h.idioma('es'); await p.wait_for_timeout(500)
        await h.shot(str(SHOTS / 'tuits_ejemplo.png'))
        await p.evaluate('go(4)'); await h.esp("document.querySelectorAll('.fa-mini').length>=8")
        check('ejemplo: 8 diapositivas (portada, 3 ideas, 3 tuits, fuentes)', await p.locator('.fa-mini').count() == 8)
        await h.shot(str(SHOTS / 'diapositivas.png'))
        # el PowerPoint del ejemplo, sin red (ni CDN)
        h.sin_cdn = True; await p.evaluate('go(5)')
        nombre = await h.descarga('#btnPptx', str(TMP / 'ejemplo.pptx')); h.sin_cdn = False
        try:
            from pptx import Presentation
            pr = Presentation(str(TMP / 'ejemplo.pptx')); txt = '\n'.join(sh.text_frame.text for s in pr.slides for sh in s.shapes if sh.has_text_frame)
            notas = [s.notes_slide.notes_text_frame.text for s in pr.slides if s.has_notes_slide]
            imgs = sum(1 for s in pr.slides for sh in s.shapes if sh.shape_type == 13)
            check('PowerPoint: se abre con python-pptx, 8 diapositivas, 3 imágenes de tuits y notas del orador', len(pr.slides) == 8 and imgs == 3 and any('demostración' in n for n in notas) and 'Fuentes' in txt, f'{len(pr.slides)} {imgs} {notas[:3]}')
            check('PowerPoint: nombre de fichero a partir del título', nombre.endswith('.pptx') and 'desinformacion' in nombre)
        except ImportError: print('  (sin python-pptx)')
        # LibreOffice lo abre y lo convierte a PDF (prueba de validez independiente)
        try:
            subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(TMP), str(TMP / 'ejemplo.pptx')], capture_output=True, timeout=180)
            check('LibreOffice convierte el PowerPoint a PDF', (TMP / 'ejemplo.pdf').exists() and (TMP / 'ejemplo.pdf').stat().st_size > 5000)
            subprocess.run(['pdftoppm', '-png', '-r', '40', '-f', '1', '-l', '3', str(TMP / 'ejemplo.pdf'), str(SHOTS / 'pptx')], capture_output=True, timeout=60)
        except (FileNotFoundError, subprocess.TimeoutExpired): print('  (sin LibreOffice)')
        # presentación a pantalla completa
        await p.click('#btnPresentar'); await p.wait_for_selector('#pres.on')
        await p.keyboard.press('ArrowRight'); await p.keyboard.press('ArrowRight')
        check('presentación: avanza con el teclado y muestra el tuit', await p.inner_text('#presNum') == '3 / 8' and await p.locator('#presEsc img').count() == 1)
        await p.keyboard.press('n'); check('presentación: N muestra las notas', await p.is_visible('#presNotas'))
        await h.shot(str(SHOTS / 'presentacion.png'), full=False)
        await p.keyboard.press('Escape'); check('presentación: Esc sale', not await p.is_visible('#pres'))

        # ── flujo real con IA simulada ──
        await p.evaluate("S=nuevoS();ENVIO_OK=false;guarda();go(1)")
        await p.evaluate("localStorage.setItem('ia_openai_key','sk-test')")
        await p.fill('#tituloPres', ''); await p.fill('#guion', GUION + '\nContacto del profesor: profe.prueba@umh.es')
        await p.click('#btnIdeasIA'); await p.wait_for_selector('#mEnvio.on', timeout=5000)
        muestra = await p.inner_text('#envioTxt')
        check('aviso previo al envío, con el correo enmascarado', '[CORREO]' in muestra and 'profe.prueba@umh.es' not in muestra)
        await p.check('#envioOk'); await p.click('#envioSi'); await h.esp("step===2 && S.ideas.length===3")
        r = await p.evaluate("S.ideas.map(i=>[i.citaOk, i.consultas.length, i.nuevos.join()])")
        check('la cita inventada de la idea 3 se marca y las búsquedas repetidas se quitan', r[0][0] and r[1][0] and not r[2][0] and r[0][1] == 2, str(r))
        check('cifra que no está en el guion (2024) señalada', r[2][2] == '2024', str(r))
        check('petición con JSON, modelo e idioma', h.seen[-1]['json'] and h.seen[-1]['model'] and 'español' in h.seen[-1]['user'])
        check('el título de la presentación lo propone la IA si estaba vacío', await p.evaluate('S.titulo') == 'Clima y opinión pública')
        check('enlaces de búsqueda a x.com con la consulta', 'x.com/search?q=consenso' in await p.inner_html('#ideas'))
        # ── añadir tuits: existe, no existe, solo en el espejo vx ──
        idea1 = '#ideas .fa-ic:nth-child(1)'
        await p.fill(f'{idea1} [data-f="pega"]', 'mira https://x.com/ciencia_abierta/status/1800000000000000111?s=20 y https://twitter.com/nadie/status/1800000000000000222 y otra vez x.com/ciencia_abierta/status/1800000000000000111')
        await p.click(f'{idea1} [data-act="tadd"]'); await h.esp("S.tuits['1800000000000000111']&&S.tuits['1800000000000000111'].estado!=='cargando'&&S.tuits['1800000000000000222'].estado!=='cargando'")
        r = await p.evaluate("[S.tuits['1800000000000000111'].estado, S.tuits['1800000000000000111'].texto, S.tuits['1800000000000000111'].avatarData.slice(0,15), S.tuits['1800000000000000111'].fotoData.length>100, S.tuits['1800000000000000222'].estado, S.tuits['1800000000000000222'].motivo, Object.keys(S.tuits).length]")
        check('tuit real: datos leídos, enlace t.co final quitado, avatar y foto descargados', r[0] == 'ok' and not r[1].endswith('abc') and r[2].startswith('data:image') and r[3], str(r))
        check('tuit inexistente: marcado como «no existe» y enlaces repetidos no se duplican', r[4] == 'error' and r[5] == 'no' and r[6] == 2, str(r))
        check('el avatar se pide en alta resolución', any('_200x200' in u for u in vistos['img']))
        idea2 = '#ideas .fa-ic:nth-child(2)'
        await p.fill(f'{idea2} [data-f="pega"]', 'https://x.com/datos_op/status/1800000000000000333')
        await p.click(f'{idea2} [data-act="tadd"]'); await h.esp("S.tuits['1800000000000000333']&&S.tuits['1800000000000000333'].estado!=='cargando'")
        check('si FxTwitter falla, se lee de vxTwitter', await p.evaluate("S.tuits['1800000000000000333'].estado==='ok' && S.tuits['1800000000000000333'].fuente==='vxTwitter'"))
        await h.shot(str(SHOTS / 'ideas.png'))
        # ── búsqueda en la web con IA (OpenAI Responses): un enlace real y otro inventado ──
        cuerpo = {}
        async def r_resp(r):
            cuerpo['openai'] = json.loads(r.request.post_data)
            txt = json.dumps({'tuits': [{'url': 'https://x.com/vecina_obs/status/1800000000000000444', 'frase': 'el clima es lo único de lo que se habla', 'porque': 'Muestra el tema en la conversación.'}, {'url': 'https://x.com/inventado/status/1800000000000000555', 'frase': 'algo', 'porque': 'x'}]})
            await r.fulfill(status=200, content_type='application/json', body=json.dumps({'output': [{'type': 'web_search_call'}, {'type': 'message', 'content': [{'type': 'output_text', 'text': txt, 'annotations': [{'type': 'url_citation', 'url': 'https://x.com/vecina_obs/status/1800000000000000444'}]}]}]}))
        await p.route('https://api.openai.com/v1/responses', r_resp)
        idea3 = '#ideas .fa-ic:nth-child(3)'
        await p.click(f'{idea3} [data-act="busca"]'); await h.esp("S.ideas[2].cand.length===2")
        r = await p.evaluate("S.ideas[2].cand.map(c=>[c.id,c.ok,c.fraseOk])")
        check('búsqueda web: el enlace inventado se detecta y la frase del real se verifica', r == [['1800000000000000444', True, True], ['1800000000000000555', False, None]] or (r[0] == ['1800000000000000444', True, True] and r[1][:2] == ['1800000000000000555', False]), str(r))
        check('búsqueda web: la petición lleva la herramienta web_search y el idioma', cuerpo['openai']['tools'][0]['type'] == 'web_search' and 'español' in cuerpo['openai']['input'])
        check('búsqueda web: el aviso del enlace inventado se ve', 'probablemente lo inventó' in await p.inner_text(idea3))
        await p.click(f'{idea3} .fa-cand[data-cid="1800000000000000444"] [data-act="cadd"]'); await h.esp("S.tuits['1800000000000000444']&&S.tuits['1800000000000000444'].estado==='ok'")
        check('añadir un candidato lo pasa a la idea con su frase clave', await p.evaluate("S.ideas[2].tuits.includes('1800000000000000444') && fraseOk(S.tuits['1800000000000000444'])"))
        # Claude: herramienta de búsqueda del servidor, limitada a x.com, y reanudación de pause_turn
        llam = []
        async def r_claude(r):
            b = json.loads(r.request.post_data); llam.append(b)
            if len(llam) == 1: return await r.fulfill(status=200, content_type='application/json', body=json.dumps({'stop_reason': 'pause_turn', 'content': [{'type': 'server_tool_use', 'id': 's1', 'name': 'web_search', 'input': {'query': 'x'}}]}))
            await r.fulfill(status=200, content_type='application/json', body=json.dumps({'stop_reason': 'end_turn', 'content': [{'type': 'web_search_tool_result', 'tool_use_id': 's1', 'content': [{'type': 'web_search_result', 'url': 'https://x.com/inventado/status/1800000000000000555', 'title': 't'}]}, {'type': 'text', 'text': '{"tuits":[]}'}]}))
        await p.route('https://api.anthropic.com/v1/messages', r_claude)
        await p.evaluate("localStorage.setItem(APP+'_prov','anthropic');localStorage.setItem(APP+'_key_anthropic','sk-ant-test');localStorage.setItem(APP+'_model_anthropic','claude-sonnet-5-5')")
        await p.click(f'{idea1} [data-act="busca"]'); await h.esp("S.ideas[0].estadoBusca && S.ideas[0].estadoBusca!=='busy'")
        t0 = llam[0]['tools'][0]
        check('Claude: web_search_20260209 limitada a x.com/twitter.com y reanuda pause_turn', t0['type'] == 'web_search_20260209' and t0['allowed_domains'] == ['x.com', 'twitter.com'] and len(llam) == 2 and llam[1]['messages'][-1]['role'] == 'assistant', json.dumps(t0))
        check('Claude: el enlace de los resultados que no existe se marca', await p.evaluate("S.ideas[0].cand.length===1 && !S.ideas[0].cand[0].ok"))
        await p.unroute('https://api.anthropic.com/v1/messages'); await p.evaluate("localStorage.setItem(APP+'_prov','openai')")
        # ── paso 3: valoración con IA ──
        await p.evaluate('go(3)'); await p.click('#btnValora'); await h.esp("S.tuits['1800000000000000111'].val && S.tuits['1800000000000000333'].val")
        r = await p.evaluate("[S.tuits['1800000000000000111'].val.fraseOk, fraseOk(S.tuits['1800000000000000111']), S.tuits['1800000000000000111'].titulo, S.tuits['1800000000000000333'].val.fraseOk, S.tuits['1800000000000000333'].frase, S.tuits['1800000000000000333'].val.nuevos.join(), !!S.tuits['1800000000000000444'].val]")
        check('valoración: frase verificada se resalta y título aplicado', r[0] and r[1] and r[2] == 'Un consenso sólido', str(r))
        check('valoración: frase que no está en el tuit se marca y no se usa', r[3] is False and r[4] == '', str(r))
        check('valoración: cifras nuevas de la IA señaladas; el id inexistente se ignora', r[5] in ('2023,65', '65,2023'), str(r))
        check('valoración: el aviso de frase no comprobada se ve', 'sin frase clave comprobada' in await p.inner_text('#estadoVal'))
        # ── capturas reales ──
        await p.evaluate("document.getElementById('capBox').open=true; localStorage.setItem(APP+'_cap_prov','microlink'); pintaCap()")
        await p.click('.fa-tw[data-tid="1800000000000000111"] [data-act="captura"]'); await h.esp("S.tuits['1800000000000000111'].captura.length>100")
        dims = await p.evaluate("cargaImg(S.tuits['1800000000000000111'].captura).then(i=>[i.naturalWidth,i.naturalHeight])")
        check('captura: pide la versión incrustada del tuit, recorta el fondo y cuenta el uso', 'platform.twitter.com%2Fembed%2FTweet.html%3Fid%3D1800000000000000111' in vistos['cap'][-1] and dims[0] < 1100 and dims[1] < 700 and await p.evaluate("usoCaptura('microlink')") == 1, str(dims))
        check('captura: el tuit pasa a usar la captura', await p.evaluate("S.tuits['1800000000000000111'].modo") == 'captura')
        await p.evaluate("localStorage.setItem(APP+'_cap_prov','screenshotone'); pintaCap()")
        await p.click('.fa-tw[data-tid="1800000000000000333"] [data-act="captura"]'); await p.wait_for_timeout(600)
        check('captura: sin la clave obligatoria no se gasta nada y se explica', len(vistos['cap']) == 1 and 'Falta la clave' in await p.inner_text('#toast'))
        await p.fill('#capKey', 'clave-prueba'); CAP['modo'] = '429'
        await p.click('.fa-tw[data-tid="1800000000000000333"] [data-act="captura"]'); await p.wait_for_timeout(1200)
        check('captura: cuota agotada (429) con mensaje claro', 'cuota gratuita' in await p.inner_text('#toast') and 'access_key=clave-prueba' in vistos['cap'][-1])
        CAP['modo'] = 'ok'
        # imagen propia para el tuit que no existe
        await p.locator('.fa-tw[data-tid="1800000000000000222"] [data-act="propia"]').click()
        await p.set_input_files('#fPropia', files=[{'name': 'mia.png', 'mimeType': 'image/png', 'buffer': png(500, 300, (240, 240, 240), (20, 20, 480, 280, (200, 30, 30)))}])
        await h.esp("S.tuits['1800000000000000222'].propia.length>100")
        check('imagen propia: un tuit no comprobable entra con tu captura', await p.evaluate("S.tuits['1800000000000000222'].modo==='propia' && enDiapos(S.tuits['1800000000000000222'])"))
        await h.shot(str(SHOTS / 'tuits.png'))
        # ── diapositivas, PowerPoint, CSV, proyecto ──
        await p.evaluate('go(4)'); await h.esp("document.querySelectorAll('.fa-mini').length>3")
        n = await p.locator('.fa-mini').count()
        await p.locator('.fa-mini').nth(1).locator('input').uncheck()
        await p.evaluate('go(5)'); await h.descarga('#btnPptx', str(TMP / 'real.pptx'))
        from pptx import Presentation
        pr = Presentation(str(TMP / 'real.pptx'))
        enl = [r.hyperlink.address for s in pr.slides for sh in s.shapes if sh.has_text_frame for pa in sh.text_frame.paragraphs for r in pa.runs if r.hyperlink.address]
        check('PowerPoint: respeta la diapositiva oculta y enlaza cada tuit original', len(pr.slides) == n - 1 and 'https://x.com/ciencia_abierta/status/1800000000000000111' in enl, f'{len(pr.slides)} vs {n} {enl}')
        await h.descarga('#btnCsv', str(TMP / 'f.csv')); csv = (TMP / 'f.csv').read_text(encoding='utf-8-sig')
        check('CSV de fuentes (BOM, «;», enlace y verificación)', csv.startswith('idea;') and 'https://x.com/ciencia_abierta/status/1800000000000000111' in csv and ';sí;' in csv)
        await h.descarga('#btnJson', str(TMP / 'p.json')); pj = (TMP / 'p.json').read_text()
        check('el proyecto exportado no contiene claves', 'sk-test' not in pj and 'clave-prueba' not in pj and 'sk-ant' not in pj and json.loads(pj)['formato'] == 'tuitia/1')
        await p.evaluate("S=nuevoS();guarda();go(5)")
        await p.set_input_files('#fJson', str(TMP / 'p.json')); await h.esp("S.ideas.length===3")
        nt = await p.evaluate("Object.keys(S.tuits).length"); check('abrir el proyecto lo recupera', nt == 4, str(nt))
        # ── persistencia, sin errores ──
        await p.wait_for_timeout(700); await p.reload(); await p.wait_for_timeout(900)
        check('persistencia tras recargar (tuits, capturas y valoración)', await p.evaluate("S.ideas.length===3 && S.tuits['1800000000000000111'].captura.length>100 && !!S.tuits['1800000000000000111'].val"))
        # sin IA: una idea por párrafo
        await p.evaluate("S=nuevoS();go(1)"); await p.fill('#guion', GUION); await p.click('#btnIdeasManual'); await h.esp('step===2')
        r = await p.evaluate("S.ideas.map(i=>[i.citaOk,i.consultas.length])")
        check('separar sin IA: una idea por párrafo, con búsquedas', len(r) == 4 and all(x[0] and x[1] >= 1 for x in r), str(r))
        # errores de la IA traducidos
        await h.idioma('pt'); await p.evaluate("go(1)")
        h.reply = lambda *a: asyncio.sleep(0, ('err', 429, '{"error":{"message":"You exceeded your current quota, billing"}}'))
        await p.click('#btnIdeasIA'); await p.wait_for_selector('#mEnvio.on'); await p.check('#envioOk'); await p.click('#envioSi'); await p.wait_for_timeout(4000)
        check('error de saldo traducido al portugués', 'não tem saldo' in await p.inner_text('#estadoGuion'), await p.inner_text('#estadoGuion'))
        await h.idioma('es')
        check('sin errores de página', not h.errors, str(h.errors))
        for l in ('es', 'en', 'pt'):
            for paso in (2, 3, 4):
                check(f'sin desborde en móvil ({l}, paso {paso})', not await h.desborde_movil(f"localStorage.setItem(APP+'_lang','{l}');initI18n();aplicaEjemplo();go({paso})"))
    print('Capturas en', SHOTS)
    sys.exit(0 if ok else 1)

if __name__ == "__main__": asyncio.run(main())
